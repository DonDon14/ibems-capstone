from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "client" / "IBEMS-System-Hierarchy.docx"

NAVY = "17324D"
BLUE = "2F6B8A"
LIGHT = "EAF1F5"
PALE = "F6F9FB"
GOLD = "D7A84B"
INK = RGBColor(23, 50, 77)
MUTED = RGBColor(82, 101, 120)
WHITE = RGBColor(255, 255, 255)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color="D5E0E8", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=120, start=160, bottom=120, end=160):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_cell_text(cell, title, subtitle, fill, title_color=WHITE):
    shade(cell, fill)
    borders(cell, color=fill)
    set_cell_margin(cell, 155, 180, 155, 180)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(12)
    r.font.color.rgb = title_color
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(subtitle)
    r2.font.name = "Aptos"
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = title_color


def connector(document, label="▼"):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(47, 107, 138)


def heading(document, text, level=1):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(4 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(18 if level == 1 else 12)
    r.font.color.rgb = INK
    return p


def add_header_footer(section, page_label):
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("IBEMS  /  CLIENT REFERENCE")
    hr.font.name = "Aptos"
    hr.font.size = Pt(8)
    hr.font.color.rgb = MUTED
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(page_label)
    fr.font.name = "Aptos"
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    add_header_footer(section, "Integrated Business Enterprise Management System")

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = MUTED
    styles["Normal"].paragraph_format.space_after = Pt(5)

    label = doc.add_paragraph()
    label.paragraph_format.space_after = Pt(6)
    lr = label.add_run("SYSTEM STRUCTURE  •  PREPARED FOR CLIENT REVIEW")
    lr.bold = True
    lr.font.name = "Aptos"
    lr.font.size = Pt(9)
    lr.font.color.rgb = RGBColor(47, 107, 138)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    tr = title.add_run("IBEMS System Hierarchy")
    tr.bold = True
    tr.font.name = "Aptos Display"
    tr.font.size = Pt(27)
    tr.font.color.rgb = INK

    sub = doc.add_paragraph("Authority, operational responsibility, and transaction accountability across the IBEMS portals.")
    sub.paragraph_format.space_after = Pt(12)
    sub.runs[0].font.size = Pt(11)

    root = doc.add_table(rows=1, cols=1)
    root.autofit = False
    root.columns[0].width = Inches(4.5)
    add_cell_text(root.cell(0, 0), "Administrator", "System governance • users • stores • oversight", NAVY)
    connector(doc)

    mid = doc.add_table(rows=1, cols=2)
    mid.autofit = False
    mid.columns[0].width = Inches(3.35)
    mid.columns[1].width = Inches(3.35)
    add_cell_text(mid.cell(0, 0), "Accounting Office", "Salary-grade profiles • debt • deductions • settlement", BLUE)
    add_cell_text(mid.cell(0, 1), "Store Supervisor", "Assigned-store review • variance governance • escalation", BLUE)
    connector(doc)

    ops = doc.add_table(rows=1, cols=2)
    ops.autofit = False
    ops.columns[0].width = Inches(3.35)
    ops.columns[1].width = Inches(3.35)
    add_cell_text(ops.cell(0, 0), "Store Officer / Cashier", "POS • inventory • receipts • day opening and closing", "4E7F73")
    add_cell_text(ops.cell(0, 1), "Employee / User", "Store catalog • purchases • personal history • balances", "4E7F73")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(7)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run("Authority flows downward. Operational records and audit evidence flow upward.")
    rr.bold = True
    rr.font.size = Pt(9.5)
    rr.font.color.rgb = INK

    note = doc.add_table(rows=1, cols=1)
    cell = note.cell(0, 0)
    shade(cell, LIGHT)
    borders(cell, "C8D9E5")
    set_cell_margin(cell, 140, 180, 140, 180)
    np = cell.paragraphs[0]
    np.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = np.add_run("Design principle")
    nr.bold = True
    nr.font.color.rgb = INK
    np.add_run("  Each portal exposes only the duties and records needed for its role, while preserving a connected audit trail.")

    doc.add_page_break()
    heading(doc, "Connected Transaction Flow")
    intro = doc.add_paragraph("A normal credit transaction moves through five accountable stages. No stage replaces or erases the source record.")
    intro.paragraph_format.space_after = Pt(10)

    steps = [
        ("1", "Profile", "Accounting assigns employment type, salary grade, compensation, and effective date. Credit is derived by policy."),
        ("2", "Purchase", "The employee selects available products; the cashier confirms identity, stock, payment method, and credit capacity."),
        ("3", "Posting", "IBEMS records the sale, product lines, payment details, inventory movement, and employee debt atomically."),
        ("4", "Review", "Store operations reconcile the business day; supervisors investigate exceptions and preserve supporting evidence."),
        ("5", "Settlement", "Accounting prepares deductions or repayments, reconciles results, and closes the period with an audit trail."),
    ]
    table = doc.add_table(rows=len(steps), cols=3)
    table.autofit = False
    widths = [Inches(0.48), Inches(1.08), Inches(5.25)]
    for row, (number, stage, detail) in zip(table.rows, steps):
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
        for cell in row.cells:
            shade(cell, PALE)
            borders(cell)
            set_cell_margin(cell, 115, 130, 115, 130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(number)
        r0.bold = True
        r0.font.size = Pt(13)
        r0.font.color.rgb = RGBColor(47, 107, 138)
        r1 = row.cells[1].paragraphs[0].add_run(stage)
        r1.bold = True
        r1.font.color.rgb = INK
        r2 = row.cells[2].paragraphs[0].add_run(detail)
        r2.font.size = Pt(8.8)
        r2.font.color.rgb = MUTED

    heading(doc, "Role Responsibility Summary", 2)
    roles = [
        ("Administrator", "Controls structure, permissions, master data, and institution-wide oversight."),
        ("Accounting Office", "Owns salary-grade profiles, dynamic credit limits, debt review, deductions, and settlement."),
        ("Store Supervisor", "Reviews assigned stores, variances, escalations, and operational compliance."),
        ("Store Officer / Cashier", "Runs POS and inventory, manages day sessions, and creates source transaction records."),
        ("Employee / User", "Views products and personal records, uses approved credit, and reviews purchases and balances."),
    ]
    rt = doc.add_table(rows=3, cols=2)
    rt.autofit = False
    for idx, (role, detail) in enumerate(roles):
        cell = rt.cell(idx // 2, idx % 2)
        shade(cell, LIGHT if idx % 2 == 0 else "F2F6F8")
        borders(cell, "D3E0E8")
        set_cell_margin(cell, 120, 150, 120, 150)
        rp = cell.paragraphs[0]
        rp.paragraph_format.space_after = Pt(2)
        r = rp.add_run(role)
        r.bold = True
        r.font.color.rgb = INK
        dp = cell.add_paragraph(detail)
        dp.paragraph_format.space_after = Pt(0)
        dp.runs[0].font.size = Pt(8.4)
    last = rt.cell(2, 1)
    shade(last, NAVY)
    borders(last, NAVY)
    lp = last.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lr = lp.add_run("ONE CONNECTED AUDIT TRAIL")
    lr.bold = True
    lr.font.color.rgb = WHITE
    lp2 = last.add_paragraph("Every role contributes evidence to the same governed transaction history.")
    lp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp2.runs[0].font.size = Pt(8.4)
    lp2.runs[0].font.color.rgb = WHITE

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
